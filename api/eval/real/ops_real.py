"""Group 3 on real inputs: Fraud Fusion, Workflow Router, Notifier, Audit Storage.

Fraud Fusion — composite cases assembled from REAL component outputs already measured in Group 1/2:
  tampering    the production-path tier/score of real receipts and ID templates (Find it again!, SIDTD;
               forged or genuine), from the tampering evaluation's cache
  face         the face agent's verdict on real photo pairs (LFW / AgeDB / CALFW / CPLFW; same or
               different person), from the face evaluation's cache
  cross_check  docrules.cross_verify_documents on real Indian names (same person written two ways, or
               another person)
A case is FRAUDULENT when it carries a forged document, a different person's face or another person's
name; CLEAN when every component comes from genuine material. Nothing here is tuned.

Workflow Router / Notifier are deterministic; real data supplies the values they handle (L&T amounts:
Head approval above Rs 5 lakh, Indian digit grouping in messages). Audit Storage stores real document
files (SROIE scans, statement PDFs, ID cards) in an isolated store and reads them back.
"""
from __future__ import annotations

import base64
import random
import re
import tempfile
from pathlib import Path

from eval import common as C
from eval.real import core as R


# ── Fraud Fusion ────────────────────────────────────────────────────────────
def _tamper_pool(split: str):
    from eval.real.evidence_real import _items, _refused
    cache = R.ItemCache("evidence", f"prod-v1-{split}")
    gen, forged = [], []
    for it in _items(split, 200):
        g = cache.get(it["id"])
        if g and g.get("tier"):
            g = _refused({**g, "id": it["id"]})    # tier / score under the current fusion (and the patch signal)
            (forged if it["forged"] else gen).append({"risk_tier": g["tier"], "risk_score": g["score"], "src": it["set"]})
    return gen, forged


def _face_pool(split: str):
    from eval.real.face_real import pairs
    cache = R.ItemCache("face_verification", "raw-v1")
    same, diff = [], []
    for p in pairs(split):
        g = cache.get(p["id"])
        if g and g.get("verdict"):
            overall = {"match": "match", "different": "mismatch"}.get(g["verdict"], "uncertain")
            face = {"overall": overall, "pairs": [{"a": "pan", "b": "aadhaar", "similarity": g["sim"], "verdict": g["verdict"]}]}
            (same if p["label"] else diff).append(face)
    return same, diff


def _names(split: str, rng: random.Random):
    from eval.real.case_real import name_pairs
    ps = name_pairs(split, 120)
    same = [p for p in ps if p["category"] in ("exact_case", "honorific", "initials", "middle_omitted", "surname_first")]
    other = [p for p in ps if p["category"] in ("different_person", "same_surname_other_person")]
    rng.shuffle(same)
    rng.shuffle(other)
    return same, other


def _checks(a: str, b: str) -> list:
    from app.docrules import cross_verify_documents
    return cross_verify_documents({"pan": {"parsed": {"full_name": a}, "label": "PAN Card"},
                                   "bank_statement": {"parsed": {"account_holder": b}, "label": "Bank Statement"}})


def eval_fraud_fusion(split: str = "test", n_clean: int = 300, n_each: int = 75) -> dict:
    from app.agents.registry import get_agent
    agent = get_agent("fraud_fusion")
    rng = random.Random(f"{R.SEED}-fusion-{split}")
    t_gen, t_forged = _tamper_pool(split)
    f_same, f_diff = _face_pool(split)
    n_same, n_other = _names(split, rng)
    if not (t_gen and t_forged and f_same and f_diff):
        raise RuntimeError("component caches are missing - run the tampering and face evaluations first")

    def case(kind: str):
        docs = {"pan": {"tampering": rng.choice(t_gen)}, "bank_statement": {"tampering": rng.choice(t_gen)}}
        face = rng.choice(f_same)
        name = rng.choice(n_same)
        if kind in ("forged_document", "forged_and_face"):
            docs[rng.choice(["pan", "bank_statement"])] = {"tampering": rng.choice(t_forged)}
        if kind in ("other_persons_face", "forged_and_face"):
            face = rng.choice(f_diff)
        if kind == "other_persons_name":
            name = rng.choice(n_other)
        return {"documents": docs, "face": face}, _checks(name["a"], name["b"])

    kinds = ["clean"] * n_clean + [k for k in ("forged_document", "other_persons_face", "other_persons_name",
                                               "forged_and_face") for _ in range(n_each)]
    rows = []
    for k in kinds:
        c, checks = case(k)
        out = agent.execute({"case": c, "cross_checks": checks, "dedupe": None, "patterns": None}).payload
        row = {"kind": k, "tier": out["tier"], "score": out["fused_score"]}
        if k == "clean":                       # monotonicity: add one forged document to this clean case
            worse = {**c, "documents": {**c["documents"], "salary_slip": {"tampering": rng.choice(t_forged)}}}
            o2 = agent.execute({"case": worse, "cross_checks": checks, "dedupe": None, "patterns": None}).payload
            row["mono_ok"] = o2["fused_score"] >= out["fused_score"] - 1e-9
        rows.append(row)
    fraud = [r for r in rows if r["kind"] != "clean"]
    clean = [r for r in rows if r["kind"] == "clean"]
    # where do flags on clean cases come from? the same clean case built from one document source at a time
    # (separate random stream: the registered rows above are not touched)
    by_source = {}
    for src in sorted({t["src"] for t in t_gen}):
        pool = [t for t in t_gen if t["src"] == src]
        r2 = random.Random(f"{R.SEED}-fusion-{split}-{src}")
        tiers = []
        for _ in range(n_clean):
            name = r2.choice(n_same)
            c = {"documents": {"pan": {"tampering": r2.choice(pool)}, "bank_statement": {"tampering": r2.choice(pool)}},
                 "face": r2.choice(f_same)}
            tiers.append(agent.execute({"case": c, "cross_checks": _checks(name["a"], name["b"]), "dedupe": None,
                                        "patterns": None}).payload["tier"])
        by_source[src] = {"clean_case_flagged": round(sum(t != "LOW" for t in tiers) / len(tiers), 4),
                          "clean_case_high": round(sum(t == "HIGH" for t in tiers) / len(tiers), 4),
                          "genuine_documents": len(pool)}
    by = {k: {"flagged": round(sum(r["tier"] in ("MEDIUM", "HIGH") for r in rs) / len(rs), 4),
              "high": round(sum(r["tier"] == "HIGH" for r in rs) / len(rs), 4), "n": len(rs)}
          for k in dict.fromkeys(r["kind"] for r in rows) for rs in [[r for r in rows if r["kind"] == k]]}
    metrics = {"fraud_case_recall": R.wilson(sum(r["tier"] in ("MEDIUM", "HIGH") for r in fraud), len(fraud)),
               "clean_case_false_high": R.wilson(sum(r["tier"] == "HIGH" for r in clean), len(clean)),
               "clean_case_flagged": R.wilson(sum(r["tier"] != "LOW" for r in clean), len(clean)),
               "monotonicity": R.wilson(sum(r["mono_ok"] for r in clean), len(clean)),
               "auc": R.auc_ci(rows, score=lambda r: r["score"], label=lambda r: r["kind"] != "clean"),
               "by_case_kind": by, "clean_cases_by_document_source": by_source}
    data = {"name": "composite cases from real component outputs: tampering (Find it again! + SIDTD, production "
                    "path), face pairs (4 benchmarks), Indian names", "split": split, "n": len(rows),
            "pools": {"tamper_genuine": len(t_gen), "tamper_forged": len(t_forged), "face_same": len(f_same),
                      "face_different": len(f_diff)}}
    return R.save("fraud_fusion", split, metrics, data)


# ── Workflow Router ─────────────────────────────────────────────────────────
def eval_workflow_router(split: str = "test", n: int = 2000) -> dict:
    from app.agents.registry import get_agent
    from eval.agents.underwriting_eval import DIAGRAM7, _expected
    from eval.real.uw_real import _loans
    agent = get_agent("workflow_router")
    ok = total = 0
    fails = []
    for product in DIAGRAM7:
        review = "underwritingReview" if product in ("msme", "working_capital") else "underReview"
        for st in ["draft", "submitted", "documentsProcessing", "documentsVerified", *DIAGRAM7[product], review,
                   "offerGenerated", "accepted", "sanctioned", "disbursed", "rejected", "cancelled"]:
            out = agent.execute({"case": {"loan": {"loan_type": product}}, "status": st}).payload
            exp = _expected(product, st)
            good = out.get("next") in exp if exp else out.get("next") is None
            total += 1
            ok += good
            if not good and len(fails) < 10:
                fails.append({"product": product, "status": st, "got": out.get("next")})
    head_ok = head_n = 0
    for x in _loans(split, n):
        amt = float(x["row"]["disbursed_amount"])
        for a in (amt, amt * 8):               # the real amount, and the same loan at a home-loan scale
            o = agent.execute({"case": {"loan": {"loan_type": "two_wheeler"}}, "status": "underReview",
                               "extra": {"loan_amount": a}}).payload
            head_n += 1
            head_ok += bool(o.get("requires_head_approval")) == (a > 500000)
    metrics = {"accuracy": R.wilson(ok, total), "head_rule_on_real_amounts": R.wilson(head_ok, head_n),
               "transitions": total}
    data = {"name": "lifecycle diagram (all products and statuses) + L&T loan amounts for the Head-approval rule",
            "split": split}
    return R.save("workflow_router", split, metrics, data, failures=fails)


# ── Notifier ────────────────────────────────────────────────────────────────
def _inr_ref(v: float) -> str:
    """Independent Indian digit grouping: 1234567 -> 12,34,567."""
    s = str(int(round(v)))
    if len(s) <= 3:
        return s
    head, tail = s[:-3], s[-3:]
    parts = []
    while len(head) > 2:
        parts.insert(0, head[-2:])
        head = head[:-2]
    if head:
        parts.insert(0, head)
    return ",".join(parts) + "," + tail


# Who must be told about each event, written from the decision workflow (officer -> Head -> borrower),
# not read from the notifier's own template table.
RECIPIENTS = {
    "application_submitted": {"officer", "user"}, "assessment_referred": {"officer"}, "underwriting_completed": {"officer"},
    "offer_drafted": {"officer"}, "more_info_requested": {"user"}, "documents_verified": {"user"},
    "documents_requested": {"user"}, "documents_uploaded": {"officer"}, "referred_to_head": {"head"},
    "sent_back": {"officer"}, "offer_made": {"user", "officer"}, "offer_accepted": {"officer"},
    "offer_declined": {"officer"}, "application_rejected": {"user"}, "loan_sanctioned": {"user"},
    "loan_disbursed": {"user"}}


def _recipients(event: str, ctx: dict) -> set:
    roles = set(RECIPIENTS[event])
    if event == "assessment_referred" and ctx.get("fraud_tier") == "HIGH":
        roles.add("head")                     # a HIGH fraud referral also goes to the Head
    if event == "underwriting_completed" and ctx.get("requires_head"):
        roles.add("head")                     # above Rs 5 lakh the Head must approve
    return roles


def eval_notifier(split: str = "test", n: int = 1000) -> dict:
    from app.agents.registry import get_agent
    from app.agents.tools import notifier
    from eval.agents.ops_eval import _FULL, _NoNetwork
    from eval.real.uw_real import _loans
    agent = get_agent("notifier")
    events = sorted({t["event"] for t in notifier.TEMPLATES.values()})
    rng = random.Random(R.SEED)
    correct = checks = 0
    fails, unknown = [], set()
    with _NoNetwork() as net:
        for x in _loans(split, n):
            amt = float(x["row"]["disbursed_amount"]) * rng.choice([1, 1, 8])
            ev = rng.choice(events)
            ctx = {**_FULL, "application_no": f"LT-{x['id']}", "amount": amt, "requires_head": amt > 500000,
                   "emi": round(amt / 24, 2), "fraud_tier": rng.choice(["LOW", "MEDIUM", "HIGH"]),
                   "actor": "Evaluation Officer", "remarks": "Please upload the last three payslips",
                   "valid_until": "15 Oct 2026", "reference": f"UTR{x['id']}"}
            if ev not in RECIPIENTS:
                unknown.add(ev)
                continue
            out = agent.execute({"case": {}, "event": ev, "context": ctx, "write": False}).payload
            drafts = [d for d in out["drafts"] if not d.get("skipped")]
            ok = {d["role"] for d in drafts} == _recipients(ev, ctx)
            for d in drafts:
                text = d["title"] + " " + d["body"]
                ok &= "{" not in text and "}" not in text and d["channel"] == "in_app"
                if "amount" in notifier.variables(notifier.TEMPLATES[d["template"]]):
                    ok &= ("Rs " + _inr_ref(amt)) in text
            checks += 1
            correct += ok
            if not ok and len(fails) < 10:
                fails.append({"event": ev, "amount": amt})
    metrics = {"template_accuracy": R.wilson(correct, checks), "external_sends": len(net.attempts), "checks": checks,
               "events_covered": len(RECIPIENTS), "events_without_expectation": sorted(unknown)}
    data = {"name": "every notification event with real L&T loan amounts (and x8 for Head-approval amounts); "
                    "amounts checked against an independent Indian digit-grouping formatter", "split": split}
    return R.save("notifier", split, metrics, data, failures=fails)


# ── Audit Storage ───────────────────────────────────────────────────────────
def eval_audit_storage(split: str = "test", n: int = 40) -> dict:
    import openpyxl
    from app import storage
    from app.agents.registry import get_agent
    from eval.datasets.synth.identities import make_aadhaar
    from eval.real.intake_real import _png, pan_items, sroie_items, statement_items
    tmp = Path(tempfile.mkdtemp(prefix="finsight-store-", dir=str(R.DATA_ROOT / "work")))
    storage.STORAGE_DIR, storage.XLSX_PATH = tmp, tmp / "applications.xlsx"
    agent = get_agent("audit_storage")
    rng = random.Random(R.SEED)
    pans, stmts, receipts = pan_items(split, n), statement_items(split, n), sroie_items(split, n)
    ok_files = total = ok_rows = 0
    secrets = []
    for k in range(n):
        p, s, r = pans[k % len(pans)], stmts[k % len(stmts)], receipts[k % len(receipts)]
        aad = make_aadhaar(rng)
        files = [("pan", p["path"].read_bytes(), "image/png", {"pan_number": p["expected"]["pan_number"],
                                                              "full_name": p["expected"]["full_name"]}),
                 ("aadhaar", _png(r["image"]), "image/png", {"aadhaar_number": f"{aad[:4]} {aad[4:8]} {aad[8:]}"}),
                 ("bank_statement", s["path"].read_bytes(), "application/pdf", {"account_number": "XXXX1234"})]
        secrets.append((p["expected"]["pan_number"], aad))
        docs = [{"doc_type": dt, "label": dt, "file_name": f"{dt}.{'pdf' if mime.endswith('pdf') else 'png'}", "mime": mime,
                 "fields": fields, "original_b64": f"data:{mime};base64," + base64.b64encode(raw).decode(),
                 "masked_b64": f"data:{mime};base64," + base64.b64encode(raw[:64]).decode(), "timings": {}}
                for dt, raw, mime, fields in files]
        rec = {"app_id": f"APP-REAL-{k:03d}", "applicant_name": "Evaluation Applicant", "loan_type": "Two Wheeler Loan",
               "documents": docs}
        res = agent.execute({"record": rec, "persist_sql": False}).payload["store"]
        ok_rows += res.get("rows") == len(docs)
        for (dt, raw, mime, _), d in zip(files, docs):
            total += 1
            found = list((tmp / rec["app_id"]).glob(f"{dt}_original.*"))
            ok_files += bool(found) and found[0].read_bytes() == raw
    wb = openpyxl.load_workbook(storage.XLSX_PATH)
    off = " ".join(str(c) for row in wb["Officer"].iter_rows(min_row=2, values_only=True) for c in row)
    mgr = " ".join(str(c) for row in wb["Manager"].iter_rows(min_row=2, values_only=True) for c in row)
    digits = lambda s: re.sub(r"\D", "", s)  # noqa: E731
    leak_o = sum(1 for pan, aad in secrets if pan in off) + sum(1 for pan, aad in secrets if aad in digits(off))
    full_m = sum(1 for pan, aad in secrets if pan in mgr)
    metrics = {"roundtrip_integrity": R.wilson(ok_files, total), "rows_ok": R.wilson(ok_rows, n),
               "officer_masking": R.wilson(2 * n - leak_o, 2 * n), "manager_sheet_full_pan": R.wilson(full_m, n),
               "n_documents": total}
    data = {"name": "real files (PAN card PNGs, SROIE receipt scans, statement PDFs) stored by the agent in an "
                    "isolated store; no SQL write", "split": split, "store": str(tmp)}
    return R.save("audit_storage", split, metrics, data)
