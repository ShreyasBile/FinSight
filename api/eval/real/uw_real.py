"""Underwriting Supervisor, Credit Memo (+ PDF / Word export) and Offer Generation on real L&T loans.

Each loan goes through what the Underwriting Flow does: Risk Scoring (real bureau and loan fields), the
evidence bundle, the deterministic baseline, the Underwriting Crew (two local-LLM calls: supervisor and
memo writer), the supervisor's floor and the memo's grounding gate. The real outcome (default) is known.

Two evidence settings, reported separately:
  as_recorded  only what L&T records about the loan: bureau data, amounts, and whether a PAN / Aadhaar
               was on file (PAN for ~8% of these loans), plus the consent the wizard now records.
               Missing KYC evidence correctly keeps most of them at "request more information".
  complete     the same loan with the evidence a completed, KYC-verified application carries (both ID
               documents read and clean, faces match) — L&T loans were real, disbursed loans, so that
               evidence existed. Only the real credit information varies. The primary metric is
               measured here.
Before 2026-10-01 neither setting could ever reach "approve": the loan-terms rules (KFS, cooling-off)
were checked at assessment, before an offer exists. They now apply from the offer stage
(compliance/regulatory_data.FROM_OFFER_STAGE) — found by this evaluation's code check.
"""
from __future__ import annotations

import io
import random
import re

from eval import common as C
from eval.real import core as R


def _complete_extra(extra: dict, rng: random.Random) -> dict:
    ex = dict(extra)
    ex.setdefault("tenure_months", rng.choice([12, 18, 24, 36]))
    return ex


def _complete_case(case: dict, rng: random.Random) -> dict:
    from eval.datasets.synth.identities import make_pan
    c = dict(case)
    c["documents"] = {
        "pan": {"identity_ok": True, "parsed": {"pan_number": make_pan(rng, "S")}, "validation": {"flags": {}},
                "tampering": {"risk_tier": "LOW", "risk_score": round(rng.uniform(0.02, 0.15), 3)}},
        "aadhaar": {"identity_ok": True, "parsed": {"aadhaar_number": "XXXX-XXXX-1234"}, "validation": {"flags": {}},
                    "tampering": {"risk_tier": "LOW", "risk_score": round(rng.uniform(0.02, 0.15), 3)}}}
    c["face"] = {"overall": "match", "pairs": [{"a": "pan", "b": "aadhaar", "similarity": round(rng.uniform(0.45, 0.8), 3),
                                                 "verdict": "match"}]}
    return c


def _policy(case, extra, assessment=None) -> str:
    from app.agents.tools import policy_context
    return policy_context.evaluate(policy_context.build_context(case, extra, assessment))["overall_status"]


def _uw_one(row, setting: str, export: bool) -> dict:
    from app.agents.base import grounded
    from app.agents.crews.underwriting import MEMO_SECTIONS, STRICTNESS, run_crew
    from app.agents.impl.underwriting import baseline_recommendation, evidence_bundle
    from app.agents.registry import get_agent
    from eval.real.risk_real import lt_inputs
    rng = random.Random(f"uw-{row['UniqueID']}")
    case, extra = lt_inputs(row)
    case = {**case, "application_no": f"LT-{row['UniqueID']}", "case_id": f"LT-{row['UniqueID']}",
            # what the wizard records at submission (all five declarations are required to submit)
            "consent": {"data_collection": True, "credit_bureau_pull": True}}
    extra = {**extra, "tenure_months": rng.choice([12, 18, 24, 36])}
    if setting == "complete":
        case, extra = _complete_case(case, rng), _complete_extra(extra, rng)
        fraud = {"tier": "LOW", "fused_score": 0.05, "floors": [], "dominant_component": None}
        case["assessment"] = {"fraud": fraud}
    else:
        case["assessment"] = {}
    case["assessment"]["policy_status"] = _policy(case, extra, case["assessment"].get("fraud"))
    risk = get_agent("risk_scoring").execute({"case": case, "income": None, "extra": extra}).payload
    ev = evidence_bundle(case, None, risk, extra)
    base = baseline_recommendation(ev)
    with C.Timer() as t:
        rec_llm, memo_llm, warns = run_crew(ev, base)
    rec = get_agent("underwriting_supervisor").execute({"evidence": ev, "llm_recommendation": rec_llm}).payload
    memo = get_agent("credit_memo").execute({"evidence": ev, "recommendation": rec, "llm_memo": memo_llm}).payload
    out = {"baseline": base["recommendation"], "llm": rec_llm.recommendation if rec_llm else None,
           "final": rec["recommendation"], "blocked": rec["loosening_blocked"],
           "never_looser": STRICTNESS[rec["recommendation"]] >= STRICTNESS[base["recommendation"]],
           "requires_head": rec["requires_head"], "rationale_grounded": rec["grounded"],
           "policy": case["assessment"]["policy_status"], "risk_advisory": risk["decision_advisory"],
           "memo_source": memo["source"], "memo_replaced": memo["sections_replaced"], "memo_complete": memo["complete"],
           "memo_final_grounded": all(grounded(memo["memo"][s], [ev, rec])[0] for s in MEMO_SECTIONS),
           "crew_s": round(t.ms / 1000, 1), "warnings": [w[:80] for w in warns][:3]}
    if memo_llm is not None:
        data = memo_llm.model_dump()
        ok = [s for s in MEMO_SECTIONS if data.get(s) and grounded(data[s], [ev, rec])[0]]
        out["llm_memo_sections_grounded"] = len(ok)
        out["llm_memo_fully_grounded"] = len(ok) == len(MEMO_SECTIONS)
    if export:
        out["export"] = export_check(case, ev, risk, rec, memo, extra, rng)
    return out


def _model(case, ev, risk, rec, memo, extra, pan, aadhaar) -> dict:
    """The export's view model, built the way memo_export.build_model builds it from SQL."""
    from datetime import datetime
    from app import memo_export as M
    body = memo["memo"]
    amount = extra.get("loan_amount")
    return {"ref": f"CM-{case['application_no']}", "generated_at": datetime.now(), "exported_at": datetime.now(),
            "prepared_by": "Evaluation", "status": "DRAFT — advisory", "source": memo["source"],
            "application_no": case["application_no"], "product": "Two Wheeler Loan", "branch": "EVAL",
            "officer": "Evaluation officer", "head": None, "applied_at": None, "status_label": "Under review",
            "applicant": {"name": "Evaluation Applicant", "pan": M.mask_pan(pan), "aadhaar": M.mask_aadhaar(aadhaar),
                          "dob": "—", "occupation": extra.get("employment_type") or "—", "location": "—"},
            "loan": {"requested": amount, "tenure": extra.get("tenure_months")},
            "recommendation": {"key": rec["recommendation"], "label": M.REC_LABEL.get(rec["recommendation"], "—"),
                               "tone": M.REC_TONE.get(rec["recommendation"], "warn"), "rationale": rec.get("rationale") or "",
                               "baseline": M.REC_LABEL.get(rec.get("baseline"), ""), "model": M.REC_LABEL.get(rec.get("llm_recommendation"), ""),
                               "requires_head": bool(rec.get("requires_head"))},
            "kpis": {"risk_score": risk.get("overall"), "risk_grade": risk.get("grade"), "fraud": "Low", "fraud_score": 0.05,
                     "policy": M.pretty(case["assessment"].get("policy_status")).capitalize(), "income_used": None,
                     "income_declared": None},
            "sections": [(title, body.get(key) or "—") for key, title in M.SECTIONS],
            "conditions": list(body.get("conditions") or []), "risk_reasons": list(risk.get("overrides") or []),
            "income": {"estimates": [], "mismatches": [], "used": None, "basis": None},
            "drivers": [{"label": M.pretty(c.get("dimension")).capitalize(), "score": c.get("score"), "weight": c.get("weight"),
                         "note": (c.get("factors") or [""])[0]} for c in (risk.get("dimension_contributions") or [])[:6]],
            "documents": [{"label": "PAN Card", "file": "pan.jpg", "status": "Verified", "tamper": "Low", "attempt": 1},
                          {"label": "Aadhaar Card", "file": "aadhaar.jpg", "status": "Verified", "tamper": "Low", "attempt": 1}],
            "offer": None, "next_step": "", "notes": list(risk.get("model_notes") or [])[:3]}


def export_check(case, ev, risk, rec, memo, extra, rng) -> dict:
    """Render the memo to PDF and Word exactly as the download endpoint does, then read both back."""
    import pypdfium2 as pdfium
    from docx import Document
    from app import memo_export as M
    from eval.datasets.synth.identities import make_aadhaar, make_pan
    pan, aadhaar = make_pan(rng, "E"), make_aadhaar(rng)
    m = _model(case, ev, risk, rec, memo, extra, pan, aadhaar)
    out = {}
    try:
        pdf = M.render_pdf(m)
        doc = pdfium.PdfDocument(pdf)
        text = " ".join(doc[i].get_textpage().get_text_range() for i in range(len(doc)))
        images = sum(1 for o in doc[0].get_objects() if o.type == pdfium.raw.FPDF_PAGEOBJ_IMAGE)
        out["pdf"] = {"opens": True, "pages": len(doc), "logo": images >= 1,
                      "ids_masked": pan not in text and aadhaar not in text and pan[5:] in text and aadhaar[-4:] in text,
                      "sections": all(t in text for _, t in M.SECTIONS[:6])}
    except Exception as exc:  # noqa: BLE001
        out["pdf"] = {"opens": False, "error": f"{type(exc).__name__}: {str(exc)[:120]}"}
    try:
        d = Document(io.BytesIO(M.render_docx(m)))
        text = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(
            c.text for t in d.tables for r in t.rows for c in r.cells)
        logo = any("image" in r.reltype for r in d.part.rels.values()) or any(
            "image" in r.reltype for s in d.sections for r in s.header.part.rels.values())
        out["docx"] = {"opens": True, "logo": logo,
                       "ids_masked": pan not in text and aadhaar not in text and pan[5:] in text and aadhaar[-4:] in text,
                       "sections": all(t in text for _, t in M.SECTIONS[:6])}
    except Exception as exc:  # noqa: BLE001
        out["docx"] = {"opens": False, "error": f"{type(exc).__name__}: {str(exc)[:120]}"}
    return out


def _loans(split: str, n: int):
    from eval.real.risk_real import lt_rows
    df = lt_rows()
    df = df[[R.split_of(i) == split for i in df["id"]]]
    ids = [x["id"] for x in R.sample([{"id": i} for i in df["id"]], n, salt="uw")]
    by = df.set_index("id")
    return [{"id": i, "row": by.loc[i]} for i in ids]


_RUNS: dict = {}


def runs(split: str = "test", n: int = 120, n_as_recorded: int = 30, n_export: int = 12):
    """Cached crew runs (the supervisor and the memo are scored from the same runs)."""
    if split in _RUNS:
        return _RUNS[split]
    C.fresh_compliance_db()
    loans = _loans(split, n)
    on_err = {"final": None, "baseline": None, "llm": None, "error": True}
    comp = R.ItemCache("underwriting", f"complete-v2-{split}").run(
        [dict(x, idx=k) for k, x in enumerate(loans)],
        lambda x: _uw_one(x["row"], "complete", x["idx"] < n_export), label="crew/complete", every=10, on_error=on_err)
    built = R.ItemCache("underwriting", f"asrecorded-v2-{split}").run(
        loans[:n_as_recorded], lambda x: _uw_one(x["row"], "as_recorded", False), label="crew/as_recorded", every=10, on_error=on_err)
    lab = {x["id"]: int(x["row"]["label"]) for x in loans}
    _RUNS[split] = ([dict(r, label=lab[x["id"]]) for x, r in zip(loans, comp)],
                    [dict(r, label=lab[x["id"]]) for x, r in zip(loans[:n_as_recorded], built)])
    return _RUNS[split]


def _rate(rs):
    return round(sum(r["label"] for r in rs) / len(rs), 4) if rs else None


def eval_supervisor(split: str = "test", n: int = 120) -> dict:
    comp, built = runs(split, n)
    ok = [r for r in comp if r.get("final")]
    approve = [r for r in ok if r["final"] == "approve_recommended"]
    base_rate = _rate(ok)

    def ratio(rs):
        a = [r for r in rs if r["final"] == "approve_recommended"]
        b = _rate(rs)
        return round(_rate(a) / b, 4) if a and b else None
    by_final = {}
    for r in ok:
        b = by_final.setdefault(r["final"], [0, 0])
        b[0] += r["label"]
        b[1] += 1
    metrics = {"approve_default_ratio": R.bootstrap(ratio, ok),
               "never_loosens": R.wilson(sum(r["never_looser"] for r in ok), len(ok)),
               "llm_ran_rate": R.wilson(sum(r["llm"] is not None for r in ok), len(ok)),
               "sample_default_rate": base_rate, "approve_default_rate": _rate(approve),
               "default_rate_by_recommendation": {k: {"default_rate": round(v[0] / v[1], 4), "n": v[1]} for k, v in by_final.items()},
               "baseline_vs_final": {f"{r['baseline']}->{r['final']}": sum(1 for x in ok if (x['baseline'], x['final']) == (r['baseline'], r['final']))
                                     for r in ok},
               "loosening_attempts_blocked": sum(bool(r["blocked"]) for r in ok),
               "rationale_grounded": R.wilson(sum(bool(r["rationale_grounded"]) for r in ok), len(ok)),
               "as_recorded": {"n": len([r for r in built if r.get("final")]),
                            "final": {k: sum(1 for r in built if r.get("final") == k) for k in
                                      ("approve_recommended", "request_more_info", "refer_to_officer", "decline_review_recommended")},
                            "policy_status": {k: sum(1 for r in built if r.get("policy") == k) for k in
                                              ("PASS", "PASS_WITH_WARNINGS", "NEEDS_INFO", "FAIL")}},
               "failed_runs": sum(1 for r in comp if r.get("error")),
               "crew_latency": R.latency([r["crew_s"] * 1000 for r in ok if r.get("crew_s")])}
    data = {"name": "L&T vehicle loans (real bureau / loan fields and real default outcome), test split",
            "split": split, "n": len(ok), "settings": {"complete": "primary", "as_recorded": "what the app supplies today"}}
    return R.save("underwriting_supervisor", split, metrics, data)


def eval_credit_memo(split: str = "test", n: int = 120) -> dict:
    comp, _ = runs(split, n)
    ok = [r for r in comp if r.get("final")]
    llm = [r for r in ok if "llm_memo_fully_grounded" in r]
    ex = [r["export"] for r in ok if r.get("export")]
    exp_ok = [all(e.get(k, {}).get(f) for k in ("pdf", "docx") for f in ("opens", "logo", "ids_masked", "sections")) for e in ex]
    replaced = {}
    for r in ok:
        for s in r["memo_replaced"]:
            replaced[s] = replaced.get(s, 0) + 1
    metrics = {"grounded_rate": R.wilson(sum(r["llm_memo_fully_grounded"] for r in llm), len(llm)),
               "section_grounded_rate": R.wilson(sum(r["llm_memo_sections_grounded"] for r in llm), 6 * len(llm)),
               "section_completeness": R.wilson(sum(bool(r["memo_complete"]) for r in ok), len(ok)),
               "final_memo_grounded": R.wilson(sum(bool(r["memo_final_grounded"]) for r in ok), len(ok)),
               "export_valid": R.wilson(sum(exp_ok), len(exp_ok)),
               "export_detail": {k: {f: sum(1 for e in ex if e.get(k, {}).get(f)) for f in ("opens", "logo", "ids_masked", "sections")}
                                 for k in ("pdf", "docx")},
               "llm_memos": len(llm), "memos": len(ok), "sections_replaced_by_template": replaced}
    data = {"name": "credit memos written for real L&T loans (complete-evidence setting); exports rendered with the "
                    "download endpoint's renderers", "split": split, "n_exports": len(ex)}
    return R.save("credit_memo", split, metrics, data)


def eval_offer_generation(split: str = "test", n: int = 2000) -> dict:
    from app.agents.impl.underwriting import emi
    from app.agents.registry import get_agent
    from eval.real.risk_real import lt_inputs
    agent = get_agent("offer_generation")
    risk_agent = get_agent("risk_scoring")
    rng = random.Random(R.SEED)
    exact = foir_ok = foir_n = 0
    fails = []
    loans = _loans(split, n)
    for x in loans:
        case, extra = lt_inputs(x["row"])
        extra["tenure_months"] = rng.choice([12, 18, 24, 36, 48])
        risk = risk_agent.execute({"case": case, "income": None, "extra": extra}).payload
        income = {"consensus_monthly": rng.choice([None, 12000.0, 20000.0, 35000.0, 60000.0])}
        out = agent.execute({"case": case, "risk": risk, "income": income, "extra": extra}).payload
        p, r_, n_ = out["offered_amount"], out["interest_rate"] / 1200, out["tenure_months"]
        ref = round(p * r_ * (1 + r_) ** n_ / ((1 + r_) ** n_ - 1), 2) if p and r_ else 0.0
        ok = abs(ref - out["emi"]) <= 0.01
        exact += ok
        if not ok and len(fails) < 5:
            fails.append({"amount": p, "rate": out["interest_rate"], "tenure": n_, "emi": out["emi"], "ref": ref})
        if income["consensus_monthly"]:
            foir_n += 1
            cap = income["consensus_monthly"] * 0.5 - float(extra.get("existing_emi") or 0)
            foir_ok += out["emi"] <= max(0.0, cap) + 0.01
        assert abs(emi(p, out["interest_rate"], n_) - out["emi"]) < 0.01 or not p
    metrics = {"emi_exact": R.wilson(exact, len(loans)), "foir_cap_respected": R.wilson(foir_ok, foir_n),
               "n": len(loans)}
    data = {"name": "L&T loan amounts (real) with the risk agent's real-input scores; tenure and income drawn "
                    "(L&T records neither)", "split": split}
    return R.save("offer_generation", split, metrics, data, failures=fails)
