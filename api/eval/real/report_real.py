"""Build the real-data evaluation report (PDF + Word, FinSight letterhead) from eval/results_real.

  python -m eval.real.report_real [--out ../docs/reports]

Numbers come only from the result files (test split unless a note says otherwise); the narrative for
each agent (what was tested, what it means, what changed) comes from report_notes.json next to the
results. Both formats are rendered from one document model, so they cannot differ.
"""
from __future__ import annotations

import argparse
import io
import json
import os
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

import yaml

from eval import common as C
from eval.real import core as R

TITLE = "FinSight — Real-Data Agent Evaluation"
SUBTITLE = "Every agent tested on its own, on real or public production-like data, against targets fixed in advance"
GROUPS = [
    ("Group 1 — real public data for the agent's own task",
     ["tampering_forensics", "authenticity", "tamper_locator", "face_verification", "liveness", "risk_scoring",
      "kb_ingestion", "regulatory_qa", "rule_authoring", "process_mining_analyst"]),
    ("Group 2 — public data for part of the task, or a lawful substitute",
     ["preprocessing", "extraction", "doc_identity", "pii_masking", "validation", "dedupe", "cross_verification",
      "fraud_patterns", "income_analysis", "policy_check", "fraud_triage", "tamper_explainer",
      "underwriting_supervisor"]),
    ("Group 3 — composite cases and deterministic agents on real inputs",
     ["fraud_fusion", "credit_memo", "offer_generation", "workflow_router", "notifier", "audit_storage"]),
]
LABEL = {"tampering_forensics": "Tampering Forensics", "authenticity": "Authenticity", "tamper_locator": "Tamper Locator",
         "face_verification": "Face Verification", "liveness": "Liveness", "risk_scoring": "Risk Scoring",
         "kb_ingestion": "KB Ingestion", "regulatory_qa": "Regulatory Q&A", "rule_authoring": "Rule Authoring",
         "process_mining_analyst": "Process-Mining Analyst", "preprocessing": "Preprocessing", "extraction": "Extraction",
         "doc_identity": "Document Identity", "pii_masking": "PII Masking", "validation": "Validation",
         "dedupe": "Dedupe", "cross_verification": "Cross Verification", "fraud_patterns": "Fraud Patterns",
         "income_analysis": "Income Analysis", "policy_check": "Policy Check", "fraud_triage": "Fraud Triage",
         "tamper_explainer": "Tamper Explainer", "underwriting_supervisor": "Underwriting Supervisor",
         "fraud_fusion": "Fraud Fusion", "credit_memo": "Credit Memo", "offer_generation": "Offer Generation",
         "workflow_router": "Workflow Router", "notifier": "Notifier", "audit_storage": "Audit Storage"}
NOT_PCT = ("auc", "f1", "skill", "ratio", "ceiling", "iou", "seconds", "violations", "sends", "mrr")


def _load(agent: str) -> dict | None:
    for split in ("test", "dev"):
        p = R.RESULTS / f"{agent}.{split}.json"
        if p.exists():
            return json.loads(p.read_text(encoding="utf-8"))
    return None


def _notes() -> dict:
    p = R.RESULTS / "report_notes.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}


def fmt(metric: str, v) -> str:
    if isinstance(v, dict):
        v = v.get("value")
    if v is None:
        return "—"
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, (int,)) and not isinstance(v, bool):
        return f"{v:,}"
    not_pct = any(k in metric for k in NOT_PCT) or metric.endswith("_s") or metric.startswith("n_")
    pct = -1 <= v <= 1 and not not_pct
    return f"{v * 100:.1f}%" if pct else (f"{v:.3f}" if abs(v) < 10 else f"{v:,.1f}")


def ci(metric: str, m) -> str:
    if not isinstance(m, dict) or not m.get("ci95") or m["ci95"][0] is None:
        return ""
    lo, hi = (float(x) for x in m["ci95"])          # 0 and 1 are stored as integers
    return f"[{fmt(metric, lo)}, {fmt(metric, hi)}]"


def target_txt(t: dict) -> str:
    if not t:
        return "—"
    return ("≥ " if t.get("higher", True) else "≤ ") + fmt(t["metric"], t["target"])


# ── document model ──────────────────────────────────────────────────────────
def _arrows(x):
    """'->' typed in the notes is shown as a proper arrow, everywhere."""
    if isinstance(x, str):
        return x.replace(" -> ", " → ")
    if isinstance(x, list):
        return [_arrows(v) for v in x]
    if isinstance(x, dict):
        return {k: _arrows(v) for k, v in x.items()}
    return x


def build() -> list:
    notes = _arrows(_notes())
    res = {a: _load(a) for _, ags in GROUPS for a in ags}
    doc = [("title", TITLE, SUBTITLE, datetime.now().strftime("%d %B %Y"))]
    ratings = [r["rating"] for r in res.values() if r and r.get("rating")]
    n_strong = ratings.count("Strong")
    doc += [("h1", "1. Summary"),
            ("kpis", [("Agents evaluated", f"{len(ratings)} / 29", "info"), ("Strong", str(n_strong), "good"),
                      ("Acceptable", str(ratings.count("Acceptable")), "warn"), ("Weak", str(ratings.count("Weak")), "bad")])]
    for para in notes.get("summary", []):
        doc.append(("p", para))
    rows = []
    for gname, ags in GROUPS:
        for a in ags:
            r = res[a]
            if not r:
                rows.append([LABEL[a], "—", "not yet run", "—", "—", "—"])
                continue
            t = r.get("target") or {}
            m = r["metrics"].get(t.get("metric"))
            rows.append([LABEL[a], (r.get("data") or {}).get("short") or notes.get(a, {}).get("dataset", ""),
                         t.get("label", ""), f"{fmt(t.get('metric', ''), r.get('primary_value'))} {ci(t.get('metric', ''), m)}".strip(),
                         target_txt(t), r.get("rating") or "—"])
    doc.append(("table", ["Agent", "Data", "Primary metric", "Result [95% CI]", "Target", "Rating"], rows,
                [0.15, 0.2, 0.27, 0.16, 0.1, 0.12], "rating"))
    if notes.get("improvements"):
        doc += [("h2", "Improvements made during the evaluation"),
                ("table", ["Agent", "Problem found on real data", "Change", "Before → after (test)"],
                 [[LABEL.get(i["agent"], i["agent"]), i["problem"], i["change"], i["effect"]] for i in notes["improvements"]],
                 [0.15, 0.32, 0.31, 0.22], None)]
    if notes.get("findings"):
        doc += [("h2", "Findings that need a decision")] + [("bullets", notes["findings"])]
    doc += [("h1", "2. How the testing was done")]
    for para in notes.get("method", []):
        doc.append(("p", para))
    if notes.get("datasets"):
        doc += [("h2", "Datasets"), ("table", ["Dataset", "What it is", "Licence / terms", "Used for"],
                                     notes["datasets"], [0.2, 0.36, 0.18, 0.26], None)]
    idle = R.RESULTS / "latency_idle.json"
    if idle.exists():
        paths = json.loads(idle.read_text(encoding="utf-8"))["paths"]
        doc += [("h2", "Latency on an idle machine"),
                ("p", "The test runs shared the PC with other evaluations, so the latencies in the agent sections are upper "
                      "bounds. These were timed afterwards on development documents, one at a time, with nothing else "
                      "running and the local model already loaded."),
                ("table", ["Path", "Documents", "Median", "Slowest"],
                 [[k, str(v["n"]), f"{v['median_s']:.1f} s", f"{v['max_s']:.1f} s"] for k, v in paths.items()],
                 [0.52, 0.16, 0.16, 0.16], None)]
    for gi, (gname, ags) in enumerate(GROUPS, start=3):
        doc += [("pagebreak",), ("h1", f"{gi}. {gname}")]
        for a in ags:
            r = res[a]
            n = notes.get(a, {})
            doc.append(("h2", LABEL[a]))
            if not r:
                doc.append(("p", "Not run yet."))
                continue
            data = r.get("data") or {}
            doc.append(("p", n.get("what") or data.get("name", "")))
            mrows = []
            t = r.get("target") or {}
            if t:
                mrows.append([t.get("label", t["metric"]), f"{fmt(t['metric'], r['primary_value'])} {ci(t['metric'], r['metrics'].get(t['metric']))}".strip(),
                              target_txt(t), r.get("rating") or "—"])
            for s in r.get("secondary") or []:
                mrows.append([s.get("label", s["metric"]), f"{fmt(s['metric'], s['value'])} {ci(s['metric'], r['metrics'].get(s['metric']))}".strip(),
                              target_txt(s), s.get("rating") or "—"])
            for k, label in (n.get("extra_metrics") or {}).items():
                v = _dig(r["metrics"], k)
                if v is not None:
                    mrows.append([label, f"{fmt(k.split('.')[-1], v)} {ci(k.split('.')[-1], v)}".strip(), "", "info"])
            doc.append(("table", ["Measure", "Result [95% CI]", "Target", "Rating"], mrows, [0.46, 0.24, 0.13, 0.17], "rating"))
            for para in n.get("findings", []):
                doc.append(("p", para))
            if n.get("limits"):
                doc.append(("note", "Limits: " + n["limits"]))
            run = r.get("run") or {}
            doc.append(("small", f"Split: {data.get('split', r.get('split'))} · n = {data.get('n') or data.get('n_pairs') or data.get('n_loans') or (data.get('n_forged_flagged', 0) + data.get('n_genuine_flagged', 0)) or '—'}"
                                 f" · model: {(run.get('llm') or {}).get('model', 'local') if isinstance(run.get('llm'), dict) else 'local'}"
                                 f" · commit {run.get('git_commit', '')}{' (+ local changes)' if run.get('git_dirty') else ''}"
                                 f" · measured {str(run.get('measured_at', ''))[:16]}"))
    doc += [("pagebreak",), ("h1", f"{len(GROUPS) + 3}. Limitations")]
    doc.append(("bullets", notes.get("limitations", [])))
    doc += [("h1", f"{len(GROUPS) + 4}. Reproducing the results")]
    for para in notes.get("reproduce", []):
        doc.append(("code", para))
    return doc


def _dig(d, dotted):
    for k in dotted.split("."):
        if not isinstance(d, dict) or k not in d:
            return None
        d = d[k]
    return d


# ── PDF ─────────────────────────────────────────────────────────────────────
def render_pdf(doc: list) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas as pdfcanvas
    from reportlab.platypus import (BaseDocTemplate, Frame, KeepTogether, PageBreak, PageTemplate, Paragraph,
                                    Spacer, Table, TableStyle)
    from app import memo_export as M
    reg, bold, ital, _ = M._register_fonts()
    c = colors.HexColor
    S = {"body": ParagraphStyle("b", fontName=reg, fontSize=9.6, leading=13.8, textColor=c(M.TEXT), spaceAfter=5),
         "small": ParagraphStyle("s", fontName=reg, fontSize=7.8, leading=10.5, textColor=c(M.MUTED), spaceAfter=6),
         "note": ParagraphStyle("n", fontName=ital, fontSize=8.8, leading=12, textColor=c(M.MUTED), spaceAfter=5),
         "code": ParagraphStyle("c", fontName="Courier", fontSize=8, leading=10.5, textColor=c(M.TEXT), spaceAfter=3),
         "h1": ParagraphStyle("h1", fontName=bold, fontSize=16, leading=20, textColor=c(M.NAVY), spaceBefore=14, spaceAfter=8,
                              keepWithNext=True),
         "h2": ParagraphStyle("h2", fontName=bold, fontSize=11.5, leading=15, textColor=c(M.DEEP), spaceBefore=10, spaceAfter=4,
                              keepWithNext=True),
         "cell": ParagraphStyle("cell", fontName=reg, fontSize=8, leading=10.4, textColor=c(M.TEXT)),
         "th": ParagraphStyle("th", fontName=bold, fontSize=7.8, leading=10, textColor=c("#FFFFFF")),
         "bullet": ParagraphStyle("bl", fontName=reg, fontSize=9.4, leading=13.2, leftIndent=12, bulletIndent=2, spaceAfter=3),
         "title": ParagraphStyle("t", fontName=bold, fontSize=26, leading=31, textColor=c(M.NAVY)),
         "sub": ParagraphStyle("st", fontName=reg, fontSize=12, leading=16, textColor=c(M.MUTED)),
         "kpi": ParagraphStyle("k", fontName=bold, fontSize=18, leading=21, textColor=c(M.NAVY)),
         "kl": ParagraphStyle("kl", fontName=bold, fontSize=7.5, leading=9.5, textColor=c(M.MUTED))}
    TONE = {"Strong": (M.GOOD, M.GOOD_BG), "Acceptable": (M.WARN, M.WARN_BG), "Weak": (M.BAD, M.BAD_BG)}
    W, H = A4
    LM = RM = 17 * mm
    BW = W - LM - RM
    P = lambda t, st="body": Paragraph(escape(str(t)).replace("\n", "<br/>"), S[st])  # noqa: E731

    class Numbered(pdfcanvas.Canvas):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self._saved = []

        def showPage(self):
            self._saved.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            n = len(self._saved)
            for st in self._saved:
                self.__dict__.update(st)
                self.saveState()
                if M.LOGO.exists():
                    lw = 36 * mm
                    iw, ih = ImageReader(str(M.LOGO)).getSize()
                    self.drawImage(str(M.LOGO), LM, H - 17 * mm, width=lw, height=lw * ih / iw, mask="auto")
                self.setFont(bold, 8); self.setFillColor(c(M.NAVY))
                self.drawRightString(W - RM, H - 11.5 * mm, "REAL-DATA AGENT EVALUATION")
                self.setStrokeColor(c(M.BLUE)); self.setLineWidth(1.1)
                self.line(LM, H - 20 * mm, W - RM, H - 20 * mm)
                self.setStrokeColor(c(M.LINE)); self.setLineWidth(0.5)
                self.line(LM, 13 * mm, W - RM, 13 * mm)
                self.setFont(reg, 7.5); self.setFillColor(c(M.MUTED))
                self.drawString(LM, 9 * mm, "FinSight AI · evaluation report · aggregate results only, no personal data")
                self.drawRightString(W - RM, 9 * mm, f"Page {self._pageNumber} of {n}")
                self.restoreState()
                super().showPage()
            super().save()

    def table(header, rows, widths, rating_col):
        data = [[P(h, "th") for h in header]] + [[P(x, "cell") for x in r] for r in rows]
        t = Table(data, colWidths=[w * BW for w in widths], repeatRows=1)
        cmds = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (0, 0), (-1, 0), c(M.DEEP)),
                ("LINEBELOW", (0, 0), (-1, -1), 0.4, c(M.LINE)),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [c("#FFFFFF"), c("#F6F8FB")]),
                ("TOPPADDING", (0, 0), (-1, -1), 3.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
                ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5)]
        if rating_col:
            for i, r in enumerate(rows, 1):
                if r[-1] in TONE:
                    fg, bg = TONE[r[-1]]
                    cmds += [("BACKGROUND", (-1, i), (-1, i), c(bg))]
        t.setStyle(TableStyle(cmds))
        return t

    story = []
    for b in doc:
        k = b[0]
        if k == "title":
            story += [Spacer(1, 40), P(b[1], "title"), Spacer(1, 6), P(b[2], "sub"), Spacer(1, 4), P(b[3], "sub"),
                      Spacer(1, 18)]
        elif k in ("h1", "h2", "p", "small", "note", "code"):
            story.append(P(b[1], "body" if k == "p" else k))
        elif k == "bullets":
            story += [Paragraph(escape(str(x)), S["bullet"], bulletText="•") for x in b[1]]
        elif k == "table":
            story += [table(b[1], b[2], b[3], b[4]), Spacer(1, 6)]
        elif k == "kpis":
            cells = [[P(v, "kpi") for _, v, _ in b[1]], [P(lbl.upper(), "kl") for lbl, _, _ in b[1]]]
            t = Table(cells, colWidths=[BW / len(b[1])] * len(b[1]))
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), c(M.TINT)), ("TOPPADDING", (0, 0), (-1, -1), 6),
                                   ("BOTTOMPADDING", (0, 0), (-1, -1), 6), ("LEFTPADDING", (0, 0), (-1, -1), 8)]))
            story += [t, Spacer(1, 8)]
        elif k == "pagebreak":
            story.append(PageBreak())
    out = io.BytesIO()
    d = BaseDocTemplate(out, pagesize=A4, leftMargin=LM, rightMargin=RM, topMargin=26 * mm, bottomMargin=19 * mm,
                        title=TITLE, author="FinSight AI")
    d.addPageTemplates([PageTemplate(frames=[Frame(LM, 19 * mm, BW, H - 45 * mm, id="f")])])
    d.build(story, canvasmaker=Numbered)
    return out.getvalue()


# ── Word ────────────────────────────────────────────────────────────────────
def render_docx(doc: list) -> bytes:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor
    from app import memo_export as M
    rgb = lambda h: RGBColor.from_string(h.lstrip("#"))  # noqa: E731
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(17)
    sec.top_margin, sec.bottom_margin = Mm(28), Mm(18)
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.color.rgb = rgb(M.TEXT)
    normal.paragraph_format.space_after = Pt(4)
    hp = sec.header.paragraphs[0]
    if M.LOGO.exists():
        hp.add_run().add_picture(str(M.LOGO), width=Mm(36))
    r = hp.add_run("\t\tREAL-DATA AGENT EVALUATION")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(8), rgb(M.NAVY)
    fp = sec.footer.paragraphs[0]
    r = fp.add_run("FinSight AI · evaluation report · aggregate results only, no personal data")
    r.font.size, r.font.color.rgb = Pt(7.5), rgb(M.MUTED)
    TONE = {"Strong": M.GOOD_BG, "Acceptable": M.WARN_BG, "Weak": M.BAD_BG}
    ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW",
             "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]

    def place(tblPr, el):
        later = ORDER[ORDER.index(el.tag.split("}")[1]) + 1:]
        for ch in tblPr:
            if ch.tag.split("}")[1] in later:
                ch.addprevious(el)
                return
        tblPr.append(el)

    def shade(cell, hex_):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_.lstrip("#"))
        cell._tc.get_or_add_tcPr().append(shd)

    def para(text, size=None, bold=False, color=None, italic=False, font=None, after=None):
        p = d.add_paragraph()
        r = p.add_run(str(text))
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = rgb(color)
        if font:
            r.font.name = font
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        return p

    for b in doc:
        k = b[0]
        if k == "title":
            para("", after=40)
            para(b[1], 24, True, M.NAVY)
            para(b[2], 12, color=M.MUTED)
            para(b[3], 11, color=M.MUTED, after=20)
        elif k == "h1":
            p = para(b[1], 15, True, M.NAVY, after=6)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.keep_with_next = True
        elif k == "h2":
            p = para(b[1], 11.5, True, M.DEEP, after=3)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.keep_with_next = True
        elif k == "p":
            para(b[1])
        elif k == "small":
            para(b[1], 7.8, color=M.MUTED, after=8)
        elif k == "note":
            para(b[1], 9, color=M.MUTED, italic=True)
        elif k == "code":
            para(b[1], 8.5, font="Consolas", after=2)
        elif k == "bullets":
            for x in b[1]:
                d.add_paragraph(str(x), style="List Bullet")
        elif k == "pagebreak":
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif k == "kpis":
            t = d.add_table(rows=2, cols=len(b[1]))
            for i, (lbl, v, _) in enumerate(b[1]):
                for row, txt, sz, bo in ((0, v, 17, True), (1, lbl.upper(), 7.5, True)):
                    cell = t.cell(row, i)
                    rr = cell.paragraphs[0].add_run(txt)
                    rr.bold, rr.font.size = bo, Pt(sz)
                    rr.font.color.rgb = rgb(M.NAVY if row == 0 else M.MUTED)
                    shade(cell, M.TINT)
            d.add_paragraph()
        elif k == "table":
            header, rows, widths, rating = b[1], b[2], b[3], b[4]
            t = d.add_table(rows=1 + len(rows), cols=len(header))
            brd = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                e = OxmlElement(f"w:{edge}")
                e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), M.LINE.lstrip("#"))
                brd.append(e)
            place(t._tbl.tblPr, brd)
            for i, h in enumerate(header):
                cell = t.cell(0, i)
                rr = cell.paragraphs[0].add_run(h)
                rr.bold, rr.font.size, rr.font.color.rgb = True, Pt(8), rgb("#FFFFFF")
                shade(cell, M.DEEP)
            for ri, row in enumerate(rows, 1):
                for ci_, v in enumerate(row):
                    cell = t.cell(ri, ci_)
                    rr = cell.paragraphs[0].add_run(str(v))
                    rr.font.size = Pt(8.5)
                    if rating and ci_ == len(row) - 1 and v in TONE:
                        shade(cell, TONE[v])
            for ci_, w in enumerate(widths):
                for cell in t.columns[ci_].cells:
                    cell.width = Mm(176 * w)
            d.add_paragraph()
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(C.API_DIR.parent / "docs" / "reports"))
    a = ap.parse_args()
    doc = build()
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    (out / "08-FinSight-Real-Data-Evaluation.pdf").write_bytes(render_pdf(doc))
    (out / "08-FinSight-Real-Data-Evaluation.docx").write_bytes(render_docx(doc))
    print("written:", out / "08-FinSight-Real-Data-Evaluation.pdf", out / "08-FinSight-Real-Data-Evaluation.docx")


if __name__ == "__main__":
    main()
