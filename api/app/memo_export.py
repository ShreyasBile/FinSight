"""The credit memo as a document: GET /lending/applications/{no}/memo?format=pdf|docx.

The Underwriting Crew drafts the memo (api/app/agents/impl/underwriting.py) and the
officer reads it on screen. This turns the same content - plus the evidence behind it
(income sources, risk drivers, document checks, the draft offer) - into a document that
can be filed, e-mailed or printed: a PDF (reportlab) or a Word file (python-docx), both
on the FinSight letterhead.

Two rules the document follows:
  * Identifiers are masked. PAN and Aadhaar appear as XXXXX1234F / XXXX-XXXX-9012, as in
    the officer's Officer sheet - a memo gets forwarded, a screen does not.
  * It says what it is: an AI-drafted, advisory document. The status reads DRAFT until a
    person has decided; the last page carries the sign-off block for that person.

Only officers and the Head may export (the memo is an internal credit document).
"""
from __future__ import annotations

import io
import os
import re
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from fastapi import APIRouter, HTTPException, Query, Response
from fastapi.concurrency import run_in_threadpool

router = APIRouter(prefix="/lending", tags=["memo"])

BRAND_DIR = Path(__file__).resolve().parents[2] / "docs" / "brand"
LOGO = BRAND_DIR / "logo-lockup.png"            # dark text on transparent: for a white page
NAVY, BLUE, DEEP = "#0E1A2B", "#0B6EF5", "#0A2A84"
TINT, LINE, TEXT, MUTED = "#EEF4FF", "#D5DDE8", "#1F2937", "#64748B"
GOOD, WARN, BAD = "#15803D", "#B45309", "#B91C1C"
GOOD_BG, WARN_BG, BAD_BG = "#E8F5EC", "#FDF3E3", "#FCE9E9"

REC_LABEL = {"approve_recommended": "Approve recommended", "request_more_info": "Request more information",
             "refer_to_officer": "Refer for a closer review", "decline_review_recommended": "Decline review recommended"}
REC_TONE = {"approve_recommended": "good", "request_more_info": "warn", "refer_to_officer": "warn",
            "decline_review_recommended": "bad"}
SECTIONS = [("executive_summary", "Executive summary"), ("applicant_profile", "Applicant profile"),
            ("income_assessment", "Income assessment"), ("risk_assessment", "Risk assessment"),
            ("fraud_and_compliance", "Fraud and compliance"), ("recommendation", "Recommendation")]
HEAD_LIMIT = 500000


# ── small formatters ────────────────────────────────────────────────────────
def inr(v, symbol="₹") -> str:
    """Indian digit grouping: 168730 -> ₹1,68,730."""
    try:
        n = int(round(float(v)))
    except (TypeError, ValueError):
        return "—"
    s, neg = str(abs(n)), n < 0
    if len(s) > 3:
        head, tail = s[:-3], s[-3:]
        head = re.sub(r"(\d)(?=(\d{2})+$)", r"\1,", head)
        s = f"{head},{tail}"
    return f"{'-' if neg else ''}{symbol}{s}"


def mask_pan(pan) -> str:
    p = str(pan or "").strip().upper()
    return "XXXXX" + p[5:] if len(p) == 10 else ("—" if not p else "XXXXX")


def mask_aadhaar(a) -> str:
    d = re.sub(r"\D", "", str(a or ""))
    return f"XXXX-XXXX-{d[-4:]}" if len(d) >= 4 else "—"


def when(iso, fmt="%d %b %Y") -> str:
    if not iso:
        return "—"
    try:
        return datetime.fromisoformat(str(iso).replace("Z", "")).strftime(fmt)
    except ValueError:
        return "—"


def pretty(s) -> str:
    return str(s or "").replace("_", " ").strip()


def words(s) -> str:
    """'documentsProcessing' -> 'documents processing'."""
    return re.sub(r"(?<!^)([A-Z])", r" \1", str(s or "")).lower()


# ── the data the document is built from ─────────────────────────────────────
def build_model(application_no: str) -> dict:
    """Everything the memo shows, gathered once, so the PDF and the Word file cannot differ."""
    from . import db
    from .agents.flows.app_evidence import application_case
    from .lending import application_detail
    app = application_detail(application_no)
    try:
        case = application_case(application_no)
    except HTTPException:
        case = {}
    memo = case.get("memo")
    if not memo:
        raise HTTPException(409, "There is no credit memo for this application yet - generate it first "
                                 "(Credit Memo tab -> Generate credit memo).")
    uw = case.get("underwriting") or {}
    risk, income, rec = uw.get("risk") or {}, uw.get("income") or {}, uw.get("recommendation") or {}
    assessment = case.get("assessment") or {}
    fraud = assessment.get("fraud") or {}
    cust = app.get("customer") or {}
    body = memo.get("memo") or {}

    report = db.query_one(
        """SELECT TOP 1 r.report_no, r.generated_at, u.full_name AS by_name
           FROM ops.generated_reports r LEFT JOIN ops.users u ON u.user_id = r.generated_by
           WHERE r.application_id = ? AND r.report_type = 'credit_memo'
           ORDER BY r.generated_at DESC""", (app["application_id"],)) or {}

    amount = app.get("loan_amount")
    offer = case.get("offer") or None
    requires_head = bool(rec.get("requires_head")) or bool(amount and float(amount) > HEAD_LIMIT)
    docs_by_type = case.get("documents") or {}
    documents = []
    for d in app.get("documents") or []:
        tier = d.get("tampering_risk_tier")
        documents.append({"label": d.get("doc_label") or "Document", "file": d.get("file_name") or "",
                          "status": pretty(d.get("verification_status")).capitalize(),
                          "tamper": tier.capitalize() if tier else "Not checked",
                          "attempt": d.get("attempt_no") or 1})
    drivers = [{"label": pretty(c.get("dimension")).capitalize(), "score": c.get("score"), "weight": c.get("weight"),
                "note": (c.get("factors") or [""])[0]} for c in (risk.get("dimension_contributions") or [])[:6]]
    return {
        "ref": report.get("report_no") or f"CM-{application_no}",
        "generated_at": report.get("generated_at"),
        "exported_at": datetime.now(),
        "prepared_by": report.get("by_name") or app.get("officer_name") or "Underwriting Officer",
        "status": "DRAFT — advisory" if memo.get("status", "draft") == "draft" else pretty(memo.get("status")).upper(),
        "source": memo.get("source") or "template",
        "application_no": application_no, "product": app.get("product_name") or "Loan",
        "branch": app.get("branch_code") or "—", "officer": app.get("officer_name") or "Unassigned",
        "head": app.get("head_name"), "applied_at": app.get("applied_at"), "status_label": app.get("status_label"),
        "applicant": {"name": app.get("customer_name") or "—", "pan": mask_pan(cust.get("pan")),
                      "aadhaar": mask_aadhaar(cust.get("aadhaar")), "dob": when(cust.get("date_of_birth")),
                      "occupation": cust.get("occupation") or "—",
                      "location": ", ".join(x for x in (cust.get("city"), cust.get("state")) if x) or "—"},
        "loan": {"requested": amount, "tenure": app.get("tenure_months")},
        "recommendation": {"key": rec.get("recommendation"), "label": REC_LABEL.get(rec.get("recommendation"), "—"),
                           "tone": REC_TONE.get(rec.get("recommendation"), "warn"), "rationale": rec.get("rationale") or "",
                           "baseline": REC_LABEL.get(rec.get("baseline"), pretty(rec.get("baseline"))),
                           "model": REC_LABEL.get(rec.get("llm_recommendation"), ""),
                           "requires_head": requires_head},
        "kpis": {"risk_score": risk.get("overall"), "risk_grade": risk.get("grade"),
                 "fraud": (fraud.get("tier") or "").capitalize() or "Not assessed",
                 "fraud_score": fraud.get("fused_score"),
                 "policy": pretty(assessment.get("policy_status")).capitalize() or "Not run",
                 "income_used": income.get("consensus_monthly"), "income_declared": income.get("declared_monthly")},
        "sections": [(title, body.get(key) or "—") for key, title in SECTIONS],
        "conditions": list(body.get("conditions") or []),
        "risk_reasons": list(risk.get("overrides") or []),
        "income": {"estimates": [{"basis": e.get("basis"), "monthly": e.get("monthly")} for e in income.get("estimates") or []],
                   "mismatches": [m.get("detail") for m in income.get("mismatches") or []],
                   "used": income.get("consensus_monthly"), "basis": income.get("consensus_basis")},
        "drivers": drivers, "documents": documents,
        "offer": ({"amount": offer.get("offered_amount"), "rate": offer.get("interest_rate"),
                   "tenure": offer.get("tenure_months"), "emi": offer.get("emi"), "fee": offer.get("processing_fee"),
                   "requested": offer.get("requested_amount"), "capped": bool(offer.get("capped_by_foir"))}
                  if offer else None),
        "next_step": (case.get("workflow") or {}).get("note") or "",
        "notes": [n for n in (risk.get("model_notes") or [])][:3],
    }


def _how(m: dict) -> str:
    """One line on how the memo was produced, shown under the recommendation."""
    return ("Drafted by FinSight AI agents running locally; every figure is checked against the case evidence. "
            if m["source"] == "llm" else
            "Generated from the case evidence by FinSight's rule-based template. ") + \
        "Advisory: this memo does not approve or reject the application."


# ── PDF ─────────────────────────────────────────────────────────────────────
def _register_fonts() -> tuple[str, str, str, str]:
    """(regular, bold, italic, rupee symbol). Arial from the system when present - it carries
    the rupee sign; otherwise Helvetica and the text 'Rs' (never a missing-glyph box)."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    fonts = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
    files = {"FS-Regular": "arial.ttf", "FS-Bold": "arialbd.ttf", "FS-Italic": "ariali.ttf", "FS-BoldItalic": "arialbi.ttf"}
    if all(os.path.exists(os.path.join(fonts, f)) for f in files.values()):
        try:
            for name, f in files.items():
                if name not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont(name, os.path.join(fonts, f)))
            pdfmetrics.registerFontFamily("FS-Regular", normal="FS-Regular", bold="FS-Bold",
                                          italic="FS-Italic", boldItalic="FS-BoldItalic")
            return "FS-Regular", "FS-Bold", "FS-Italic", "₹"
        except Exception:  # noqa: BLE001 - fall back below
            pass
    return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Rs "


def render_pdf(m: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas as pdfcanvas
    from reportlab.platypus import (BaseDocTemplate, Frame, KeepTogether, PageTemplate, Paragraph, Spacer, Table,
                                    TableStyle)

    reg, bold, ital, rs = _register_fonts()
    c = colors.HexColor
    money = lambda v: inr(v, rs)  # noqa: E731

    def style(name, **kw):
        base = dict(fontName=reg, fontSize=9.5, leading=13.5, textColor=c(TEXT), alignment=TA_LEFT)
        base.update(kw)
        return ParagraphStyle(name, **base)

    S = {"body": style("body"), "small": style("small", fontSize=8, leading=11, textColor=c(MUTED)),
         "h1": style("h1", fontName=bold, fontSize=22, leading=26, textColor=c(NAVY)),
         "sub": style("sub", fontSize=10.5, leading=14, textColor=c(MUTED)),
         "h2": style("h2", fontName=bold, fontSize=11.5, leading=15, textColor=c(DEEP), spaceBefore=12, spaceAfter=4),
         "label": style("label", fontName=bold, fontSize=7.5, leading=10, textColor=c(MUTED)),
         "val": style("val", fontSize=9.5, leading=12.5),
         "kpi": style("kpi", fontName=bold, fontSize=17, leading=20, textColor=c(NAVY)),
         "cell": style("cell", fontSize=8.8, leading=11.5), "cellb": style("cellb", fontName=bold, fontSize=8.8, leading=11.5),
         "th": style("th", fontName=bold, fontSize=8, leading=10.5, textColor=c("#FFFFFF")),
         "bullet": style("bullet", leftIndent=12, bulletIndent=2, spaceAfter=2, bulletFontName=reg)}

    def P(text, st="body"):
        return Paragraph(escape(str(text)).replace("\n", "<br/>"), S[st])

    def tone_colors(tone):
        return {"good": (GOOD, GOOD_BG), "warn": (WARN, WARN_BG), "bad": (BAD, BAD_BG)}.get(tone, (MUTED, TINT))

    def grid(rows, widths, header=True, zebra=True):
        t = Table(rows, colWidths=widths, repeatRows=1 if header else 0)
        cmds = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, -1), 0.4, c(LINE)),
                ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6)]
        if header:
            cmds += [("BACKGROUND", (0, 0), (-1, 0), c(DEEP))]
        if zebra:
            cmds += [("ROWBACKGROUNDS", (0, 1 if header else 0), (-1, -1), [c("#FFFFFF"), c("#F6F8FB")])]
        t.setStyle(TableStyle(cmds))
        return t

    W, H = A4
    LM = RM = 18 * mm
    body_w = W - LM - RM
    out = io.BytesIO()

    class Numbered(pdfcanvas.Canvas):
        """Two passes' worth of state in one: 'Page X of Y' needs the total."""
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self._saved = []

        def showPage(self):
            self._saved.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            n = len(self._saved)
            for state in self._saved:
                self.__dict__.update(state)
                self._chrome(n)
                super().showPage()
            super().save()

        def _chrome(self, total):
            self.saveState()
            # letterhead
            if LOGO.exists():
                lw = 40 * mm
                ih, iw = ImageReader(str(LOGO)).getSize()[1], ImageReader(str(LOGO)).getSize()[0]
                self.drawImage(str(LOGO), LM, H - 18 * mm, width=lw, height=lw * ih / iw, mask="auto")
            self.setFont(bold, 8.5); self.setFillColor(c(NAVY))
            self.drawRightString(W - RM, H - 11.5 * mm, "CREDIT MEMORANDUM")
            self.setFont(reg, 8); self.setFillColor(c(MUTED))
            self.drawRightString(W - RM, H - 15.5 * mm, f"{m['ref']}  ·  {m['application_no']}")
            self.setStrokeColor(c(BLUE)); self.setLineWidth(1.2)
            self.line(LM, H - 21 * mm, W - RM, H - 21 * mm)
            # footer
            self.setStrokeColor(c(LINE)); self.setLineWidth(0.5)
            self.line(LM, 14 * mm, W - RM, 14 * mm)
            self.setFont(reg, 7.5); self.setFillColor(c(MUTED))
            self.drawString(LM, 10 * mm, "FinSight AI  ·  Confidential — internal credit use only  ·  "
                                         "AI-assisted draft: advisory, a person decides")
            self.drawRightString(W - RM, 10 * mm, f"Page {self._pageNumber} of {total}")
            self.restoreState()

    story = []
    # ── title block
    story += [Spacer(1, 2), P("Credit Memorandum", "h1"),
              P(f"{m['product']}  ·  {m['application_no']}", "sub"), Spacer(1, 8)]
    rec = m["recommendation"]
    fg, bg = tone_colors(rec["tone"])
    status_tbl = Table([[Paragraph(f'<font color="{fg}"><b>{escape(rec["label"])}</b></font>',
                                   style("rb", fontSize=10.5, leading=14)),
                         Paragraph(f'<font color="{MUTED}">Status</font>  <b>{escape(m["status"])}</b>',
                                   style("st", fontSize=9, leading=12)),
                         Paragraph(f'<font color="{MUTED}">Memo date</font>  <b>{when(m["generated_at"] or m["exported_at"])}</b>',
                                   style("dt", fontSize=9, leading=12))]],
                       colWidths=[body_w * 0.44, body_w * 0.30, body_w * 0.26])
    status_tbl.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, 0), c(bg)), ("BOX", (0, 0), (0, 0), 0.6, c(fg)),
                                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 9),
                                    ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
    story += [status_tbl, Spacer(1, 10)]

    # ── the case at a glance
    a, ln = m["applicant"], m["loan"]
    kv = lambda k, v: [Paragraph(escape(k.upper()), S["label"]), Paragraph(escape(str(v)), S["val"])]  # noqa: E731
    glance = [kv("Applicant", a["name"]) + kv("Application", m["application_no"]),
              kv("PAN", a["pan"]) + kv("Loan product", m["product"]),
              kv("Aadhaar", a["aadhaar"]) + kv("Amount requested", money(ln["requested"])),
              kv("Date of birth", a["dob"]) + kv("Tenure", f"{ln['tenure']} months" if ln["tenure"] else "—"),
              kv("Occupation", a["occupation"]) + kv("Branch / officer", f"{m['branch']}  ·  {m['officer']}"),
              kv("Location", a["location"]) + kv("Applied on", when(m["applied_at"]))]
    gt = Table(glance, colWidths=[body_w * 0.17, body_w * 0.32, body_w * 0.20, body_w * 0.31])
    gt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), c(TINT)), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                            ("LEFTPADDING", (0, 0), (-1, -1), 8), ("BOX", (0, 0), (-1, -1), 0.5, c(LINE))]))
    story += [gt, Spacer(1, 10)]

    # ── KPI strip
    k = m["kpis"]
    risk_tone = "bad" if (k["risk_score"] or 0) >= 70 else "warn" if (k["risk_score"] or 0) >= 40 else "good"
    fraud_tone = {"High": "bad", "Medium": "warn", "Low": "good", "None": "good"}.get(k["fraud"], "warn")
    pol = k["policy"].lower()
    pol_tone = "bad" if "fail" in pol else "good" if pol in ("pass", "passed") else "warn"
    cells = [("RISK SCORE", f"{k['risk_score']:.0f} / 100" if k["risk_score"] is not None else "—",
              f"Grade {k['risk_grade']}" if k["risk_grade"] else "Not scored", risk_tone),
             ("FRAUD RISK", k["fraud"], f"score {k['fraud_score']:.2f}" if k["fraud_score"] is not None else "", fraud_tone),
             ("POLICY CHECK", k["policy"].replace("_", " "), "RBI and internal rules", pol_tone),
             ("INCOME USED", money(k["income_used"]) if k["income_used"] else "—", "per month" +
              (f" (declared {money(k['income_declared'])})" if k["income_declared"] and k["income_declared"] != k["income_used"] else ""),
              "warn" if (k["income_declared"] and k["income_declared"] != k["income_used"]) else "good")]
    kcells = []
    for label, val, sub, tone in cells:
        col, _ = tone_colors(tone)
        kcells.append([Paragraph(label, S["label"]),
                       Paragraph(f'<font color="{col}">{escape(str(val))}</font>', S["kpi"]),
                       Paragraph(escape(sub), S["small"])])
    kt = Table([kcells], colWidths=[body_w / 4] * 4)
    kt.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, c(LINE)), ("INNERGRID", (0, 0), (-1, -1), 0.5, c(LINE)),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 7),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 7), ("LEFTPADDING", (0, 0), (-1, -1), 9)]))
    story += [kt]

    # ── recommendation
    rat = [P(rec["rationale"] or "—")]
    meta = f"Rule baseline: {rec['baseline'] or '—'}" + (f"  ·  model suggested: {rec['model']}" if rec["model"] else "")
    rat += [Spacer(1, 3), P(meta + "  ·  " + _how(m), "small")]
    if rec["requires_head"]:
        rat += [Spacer(1, 3), Paragraph(f'<font color="{WARN}"><b>Head Manager approval is required</b></font> '
                                        f'(amount above {money(HEAD_LIMIT)}).', S["small"])]
    story += [Paragraph("Recommendation", S["h2"])] + rat

    # ── narrative sections (the memo body), each with its evidence where there is some
    n = 0
    for title, text in m["sections"]:
        if title == "Recommendation":
            continue                                 # already shown above, with the rationale
        n += 1
        block = [Paragraph(f"{n}.  {escape(title)}", S["h2"]), P(text)]
        if title == "Income assessment":
            inc = m["income"]
            if inc["estimates"]:
                rows = [[Paragraph("INCOME SOURCE", S["th"]), Paragraph("PER MONTH", S["th"])]]
                rows += [[P(e["basis"], "cell"), P(money(e["monthly"]), "cell")] for e in inc["estimates"]]
                block += [Spacer(1, 5), grid(rows, [body_w * 0.72, body_w * 0.28])]
            for mm_ in inc["mismatches"]:
                block += [Spacer(1, 2), Paragraph(f'<font color="{WARN}">▲</font>  {escape(str(mm_))}', S["small"])]
            if inc["used"]:
                block += [Spacer(1, 3), P(f"Used for underwriting: {money(inc['used'])} per month"
                                          + (f" ({inc['basis']})" if inc["basis"] else ""), "small")]
        if title == "Risk assessment" and m["drivers"]:
            rows = [[Paragraph(x, S["th"]) for x in ("RISK DIMENSION", "SCORE", "WEIGHT", "MAIN FACTOR")]]
            rows += [[P(d["label"], "cellb"), P(f"{d['score']:.0f}" if d["score"] is not None else "—", "cell"),
                      P(f"{d['weight']:.0%}" if d["weight"] is not None else "—", "cell"), P(d["note"], "cell")]
                     for d in m["drivers"]]
            block += [Spacer(1, 5), grid(rows, [body_w * 0.24, body_w * 0.10, body_w * 0.10, body_w * 0.56])]
            for r in m["risk_reasons"]:
                block += [Spacer(1, 2), Paragraph(f'<font color="{BAD}">■</font>  {escape(str(r))}', S["small"])]
        if title == "Fraud and compliance" and m["documents"]:
            rows = [[Paragraph(x, S["th"]) for x in ("DOCUMENT", "VERIFICATION", "TAMPERING CHECK", "UPLOAD")]]
            rows += [[P(d["label"], "cellb"), P(d["status"], "cell"), P(d["tamper"], "cell"),
                      P("original" if d["attempt"] == 1 else f"attempt {d['attempt']}", "cell")] for d in m["documents"]]
            block += [Spacer(1, 5), grid(rows, [body_w * 0.34, body_w * 0.22, body_w * 0.26, body_w * 0.18])]
        story += [KeepTogether(block[:2])] + block[2:]

    # ── conditions and terms
    if m["conditions"]:
        story += [Paragraph("Conditions before any approval", S["h2"])]
        story += [Paragraph(escape(str(x)), S["bullet"], bulletText="•") for x in m["conditions"]]
    o = m["offer"]
    if o:
        rows = [[Paragraph(x, S["th"]) for x in ("AMOUNT OFFERED", "INTEREST RATE", "TENURE", "EMI", "PROCESSING FEE")],
                [P(money(o["amount"]), "cellb"), P(f"{o['rate']}% p.a." if o["rate"] is not None else "—", "cellb"),
                 P(f"{o['tenure']} months", "cellb"), P(money(o["emi"]), "cellb"), P(money(o["fee"]), "cellb")]]
        note = (f"Draft. Reduced from the {money(o['requested'])} requested so that total EMIs stay within 50% of income."
                if o["capped"] else "Draft offer prepared by the Offer Generation agent. It is not sent to the borrower "
                                    "until an officer or the Head approves it.")
        story += [KeepTogether([Paragraph("Proposed terms", S["h2"]), grid(rows, [body_w * 0.2] * 5, zebra=False),
                                Spacer(1, 3), P(note, "small")])]
    if m["next_step"]:
        story += [Paragraph("Next step", S["h2"]), P(m["next_step"])]

    # ── sign-off
    sign = Table([[Paragraph("PREPARED BY", S["label"]), "", Paragraph("REVIEWED / APPROVED BY", S["label"])],
                  [P(f"{m['prepared_by']}\nUnderwriting Officer", "val"), "",
                   P(f"{m['head'] or 'Head Manager'}\nHead Manager" if rec["requires_head"] else "Reviewing authority", "val")],
                  ["", "", ""],
                  [Paragraph("Signature and date", S["small"]), "", Paragraph("Signature and date", S["small"])]],
                 colWidths=[body_w * 0.44, body_w * 0.12, body_w * 0.44], rowHeights=[11, 27, 18, 12])
    sign.setStyle(TableStyle([("LINEABOVE", (0, 3), (0, 3), 0.6, c(MUTED)), ("LINEABOVE", (2, 3), (2, 3), 0.6, c(MUTED)),
                              ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story += [Spacer(1, 4), KeepTogether([Paragraph("Sign-off", S["h2"]), sign])]

    doc = BaseDocTemplate(out, pagesize=A4, leftMargin=LM, rightMargin=RM, topMargin=27 * mm, bottomMargin=20 * mm,
                          title=f"Credit Memorandum {m['application_no']}", author="FinSight AI",
                          subject=f"{m['product']} — {m['application_no']}")
    doc.addPageTemplates([PageTemplate(id="memo", frames=[Frame(LM, 20 * mm, body_w, H - 47 * mm, id="f",
                                                                leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)])])
    doc.build(story, canvasmaker=Numbered)
    return out.getvalue()


# ── Word ────────────────────────────────────────────────────────────────────
def render_docx(m: dict) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    money = lambda v: inr(v, "₹")  # noqa: E731
    rgb = lambda h: RGBColor.from_string(h.lstrip("#"))  # noqa: E731
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(18)
    sec.top_margin, sec.bottom_margin = Mm(30), Mm(20)
    sec.header_distance, sec.footer_distance = Mm(8), Mm(8)
    content_w = Mm(210 - 36)

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.color.rgb = rgb(TEXT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12

    def shade(cell, hex_):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_.lstrip("#"))
        tcPr.append(shd)

    # w:tblPr children must appear in schema order, or Word may call the file corrupt
    TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
                   "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar",
                   "tblLook", "tblCaption", "tblDescription"]

    def place(tblPr, el):
        name = el.tag.split("}")[1]
        later = TBLPR_ORDER[TBLPR_ORDER.index(name) + 1:]
        for child in tblPr:
            if child.tag.split("}")[1] in later:
                child.addprevious(el)
                return
        tblPr.append(el)

    def borders(table, inside=True, color=LINE):
        b = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right") + (("insideH", "insideV") if inside else ()):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), color.lstrip("#"))
            b.append(e)
        place(table._tbl.tblPr, b)

    def cell_pad(table, top=50, bottom=50, left=100, right=100):
        mar = OxmlElement("w:tblCellMar")
        for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa")
            mar.append(e)
        place(table._tbl.tblPr, mar)

    def text(par, s, size=None, bold=False, color=None, italic=False):
        r = par.add_run(str(s))
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = rgb(color)
        return r

    def para(s="", size=None, bold=False, color=None, after=None, before=None, align=None, italic=False):
        p = doc.add_paragraph()
        if s != "":
            text(p, s, size, bold, color, italic)
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        if before is not None:
            p.paragraph_format.space_before = Pt(before)
        if align:
            p.alignment = align
        return p

    def heading(s):
        p = para(s, 12, True, DEEP, after=3, before=11)
        p.paragraph_format.keep_with_next = True
        return p

    def set_widths(table, widths_mm):
        table.autofit = False
        for row in table.rows:
            for cell, w in zip(row.cells, widths_mm):
                cell.width = Mm(w)

    def data_table(headers, rows, widths_mm):
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders(t); cell_pad(t)
        for c_, h in zip(t.rows[0].cells, headers):
            shade(c_, DEEP)
            c_.paragraphs[0].paragraph_format.space_after = Pt(0)
            text(c_.paragraphs[0], h, 8, True, "#FFFFFF")
        for i, r in enumerate(rows):
            cells = t.add_row().cells
            for j, (c_, v) in enumerate(zip(cells, r)):
                c_.paragraphs[0].paragraph_format.space_after = Pt(0)
                text(c_.paragraphs[0], v, 9, bold=(j == 0 and len(headers) > 2))
                if i % 2 == 1:
                    shade(c_, "#F6F8FB")
        set_widths(t, widths_mm)
        # repeat the header row on each page
        trPr = t.rows[0]._tr.get_or_add_trPr()
        h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
        return t

    # ── letterhead (header) and footer
    hdr = sec.header
    ht = hdr.add_table(rows=1, cols=2, width=content_w)
    hdr.paragraphs[0]._element.getparent().remove(hdr.paragraphs[0]._element)
    left, right = ht.rows[0].cells
    if LOGO.exists():
        left.paragraphs[0].add_run().add_picture(str(LOGO), width=Mm(40))
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    text(right.paragraphs[0], "CREDIT MEMORANDUM", 8.5, True, NAVY)
    rp = right.add_paragraph(); rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    text(rp, f"{m['ref']}  ·  {m['application_no']}", 8, color=MUTED)
    set_widths(ht, [90, 84])
    rule = hdr.add_paragraph()
    pPr = rule._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "10"); bt.set(qn("w:space"), "1"); bt.set(qn("w:color"), BLUE.lstrip("#"))
    pb.append(bt); pPr.append(pb)
    rule.paragraph_format.space_after = Pt(0)

    fp = sec.footer.paragraphs[0]
    text(fp, "FinSight AI  ·  Confidential — internal credit use only  ·  AI-assisted draft: advisory, a person decides"
             "        Page ", 7.5, color=MUTED)
    for kind in ("PAGE", " of ", "NUMPAGES"):
        if kind.strip() in ("PAGE", "NUMPAGES"):
            r = fp.add_run(); r.font.size = Pt(7.5); r.font.color.rgb = rgb(MUTED)
            for t_, extra in (("begin", None), (None, kind), ("end", None)):
                if t_:
                    f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), t_); r._r.append(f)
                else:
                    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = extra; r._r.append(it)
        else:
            text(fp, kind, 7.5, color=MUTED)

    # ── title block
    para("Credit Memorandum", 24, True, NAVY, after=0)
    para(f"{m['product']}  ·  {m['application_no']}", 11, color=MUTED, after=8)
    rec = m["recommendation"]
    fg, bg = {"good": (GOOD, GOOD_BG), "warn": (WARN, WARN_BG), "bad": (BAD, BAD_BG)}.get(rec["tone"], (MUTED, TINT))
    st = doc.add_table(rows=1, cols=3); borders(st, inside=False, color=LINE); cell_pad(st, 90, 90, 140, 140)
    a, b_, c_ = st.rows[0].cells
    shade(a, bg)
    a.paragraphs[0].paragraph_format.space_after = Pt(0); text(a.paragraphs[0], rec["label"], 11, True, fg)
    b_.paragraphs[0].paragraph_format.space_after = Pt(0)
    text(b_.paragraphs[0], "Status  ", 9, color=MUTED); text(b_.paragraphs[0], m["status"], 9, True)
    c_.paragraphs[0].paragraph_format.space_after = Pt(0)
    text(c_.paragraphs[0], "Memo date  ", 9, color=MUTED)
    text(c_.paragraphs[0], when(m["generated_at"] or m["exported_at"]), 9, True)
    set_widths(st, [77, 52, 45])
    para("", after=4)

    # ── case at a glance
    ap, ln = m["applicant"], m["loan"]
    pairs = [(("Applicant", ap["name"]), ("Application", m["application_no"])),
             (("PAN", ap["pan"]), ("Loan product", m["product"])),
             (("Aadhaar", ap["aadhaar"]), ("Amount requested", money(ln["requested"]))),
             (("Date of birth", ap["dob"]), ("Tenure", f"{ln['tenure']} months" if ln["tenure"] else "—")),
             (("Occupation", ap["occupation"]), ("Branch / officer", f"{m['branch']}  ·  {m['officer']}")),
             (("Location", ap["location"]), ("Applied on", when(m["applied_at"])))]
    gt = doc.add_table(rows=len(pairs), cols=4); borders(gt, inside=False); cell_pad(gt, 45, 45, 120, 120)
    for row, ((k1, v1), (k2, v2)) in zip(gt.rows, pairs):
        for cell, (txt, isk) in zip(row.cells, ((k1.upper(), True), (v1, False), (k2.upper(), True), (v2, False))):
            shade(cell, TINT)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            text(cell.paragraphs[0], txt, 7.5 if isk else 9.5, bold=isk, color=MUTED if isk else None)
    set_widths(gt, [29, 56, 35, 54])
    para("", after=4)

    # ── KPI strip
    k = m["kpis"]
    tone = lambda t: {"good": GOOD, "warn": WARN, "bad": BAD}[t]  # noqa: E731
    rt = "bad" if (k["risk_score"] or 0) >= 70 else "warn" if (k["risk_score"] or 0) >= 40 else "good"
    ft = {"High": "bad", "Medium": "warn", "Low": "good", "None": "good"}.get(k["fraud"], "warn")
    pol = k["policy"].lower()
    pt = "bad" if "fail" in pol else "good" if pol in ("pass", "passed") else "warn"
    mism = bool(k["income_declared"] and k["income_declared"] != k["income_used"])
    kp = [("RISK SCORE", f"{k['risk_score']:.0f} / 100" if k["risk_score"] is not None else "—",
           f"Grade {k['risk_grade']}" if k["risk_grade"] else "Not scored", rt),
          ("FRAUD RISK", k["fraud"], f"score {k['fraud_score']:.2f}" if k["fraud_score"] is not None else "", ft),
          ("POLICY CHECK", k["policy"].replace("_", " "), "RBI and internal rules", pt),
          ("INCOME USED", money(k["income_used"]) if k["income_used"] else "—",
           "per month" + (f" (declared {money(k['income_declared'])})" if mism else ""), "warn" if mism else "good")]
    kt = doc.add_table(rows=1, cols=4); borders(kt); cell_pad(kt, 80, 80, 120, 100)
    for cell, (lab, val, sub, tn) in zip(kt.rows[0].cells, kp):
        cell.paragraphs[0].paragraph_format.space_after = Pt(0); text(cell.paragraphs[0], lab, 7.5, True, MUTED)
        p2 = cell.add_paragraph(); p2.paragraph_format.space_after = Pt(0); text(p2, val, 15, True, tone(tn))
        p3 = cell.add_paragraph(); p3.paragraph_format.space_after = Pt(0); text(p3, sub, 8, color=MUTED)
    set_widths(kt, [43.5] * 4)

    # ── recommendation
    heading("Recommendation")
    para(rec["rationale"] or "—")
    para(f"Rule baseline: {rec['baseline'] or '—'}" + (f"  ·  model suggested: {rec['model']}" if rec["model"] else "")
         + "  ·  " + _how(m), 8, color=MUTED, after=2)
    if rec["requires_head"]:
        p = para("", 8, after=2)
        text(p, "Head Manager approval is required ", 8, True, WARN)
        text(p, f"(amount above {money(HEAD_LIMIT)}).", 8, color=MUTED)

    # ── narrative sections with their evidence
    n = 0
    for title, body in m["sections"]:
        if title == "Recommendation":
            continue
        n += 1
        heading(f"{n}.  {title}")
        para(body)
        if title == "Income assessment":
            inc = m["income"]
            if inc["estimates"]:
                data_table(["INCOME SOURCE", "PER MONTH"], [[e["basis"], money(e["monthly"])] for e in inc["estimates"]], [125, 49])
            for x in inc["mismatches"]:
                p = para("", 8, after=1, before=2); text(p, "▲  ", 8, color=WARN); text(p, x, 8, color=MUTED)
            if inc["used"]:
                para(f"Used for underwriting: {money(inc['used'])} per month" + (f" ({inc['basis']})" if inc["basis"] else ""),
                     8, color=MUTED, before=3)
        if title == "Risk assessment" and m["drivers"]:
            data_table(["RISK DIMENSION", "SCORE", "WEIGHT", "MAIN FACTOR"],
                       [[d["label"], f"{d['score']:.0f}" if d["score"] is not None else "—",
                         f"{d['weight']:.0%}" if d["weight"] is not None else "—", d["note"]] for d in m["drivers"]],
                       [42, 17, 17, 98])
            for x in m["risk_reasons"]:
                p = para("", 8, after=1, before=2); text(p, "■  ", 8, color=BAD); text(p, x, 8, color=MUTED)
        if title == "Fraud and compliance" and m["documents"]:
            data_table(["DOCUMENT", "VERIFICATION", "TAMPERING CHECK", "UPLOAD"],
                       [[d["label"], d["status"], d["tamper"], "original" if d["attempt"] == 1 else f"attempt {d['attempt']}"]
                        for d in m["documents"]], [59, 38, 45, 32])

    if m["conditions"]:
        heading("Conditions before any approval")
        for x in m["conditions"]:
            p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(2); text(p, x)
    o = m["offer"]
    if o:
        heading("Proposed terms")
        data_table(["AMOUNT OFFERED", "INTEREST RATE", "TENURE", "EMI", "PROCESSING FEE"],
                   [[money(o["amount"]), f"{o['rate']}% p.a." if o["rate"] is not None else "—",
                     f"{o['tenure']} months", money(o["emi"]), money(o["fee"])]], [34.8] * 5)
        para((f"Draft. Reduced from the {money(o['requested'])} requested so that total EMIs stay within 50% of income."
              if o["capped"] else "Draft offer prepared by the Offer Generation agent. It is not sent to the borrower "
                                  "until an officer or the Head approves it."), 8, color=MUTED, before=3)
    if m["next_step"]:
        heading("Next step")
        para(m["next_step"])

    # ── sign-off
    heading("Sign-off")
    sg = doc.add_table(rows=3, cols=2); cell_pad(sg, 40, 40, 0, 120)
    labels = ("PREPARED BY", "REVIEWED / APPROVED BY")
    names = (f"{m['prepared_by']}\nUnderwriting Officer",
             f"{m['head'] or 'Head Manager'}\nHead Manager" if rec["requires_head"] else "Reviewing authority")
    for i in range(2):
        sg.rows[0].cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
        text(sg.rows[0].cells[i].paragraphs[0], labels[i], 7.5, True, MUTED)
        sg.rows[1].cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
        text(sg.rows[1].cells[i].paragraphs[0], names[i], 9.5)
        c2 = sg.rows[2].cells[i]; c2.paragraphs[0].paragraph_format.space_before = Pt(26)
        c2.paragraphs[0].paragraph_format.space_after = Pt(0)
        text(c2.paragraphs[0], "Signature and date", 8, color=MUTED)
        tcPr = c2._tc.get_or_add_tcPr(); tb = OxmlElement("w:tcBorders"); tp = OxmlElement("w:top")
        tp.set(qn("w:val"), "single"); tp.set(qn("w:sz"), "6"); tp.set(qn("w:color"), MUTED.lstrip("#")); tb.append(tp); tcPr.append(tb)
    set_widths(sg, [85, 85])

    core = doc.core_properties
    core.title, core.author, core.subject = f"Credit Memorandum {m['application_no']}", "FinSight AI", f"{m['product']} — {m['application_no']}"
    zoom = doc.settings.element.find(qn("w:zoom"))      # python-docx's template omits the required percent
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── the endpoint ────────────────────────────────────────────────────────────
MEDIA = {"pdf": "application/pdf",
         "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}


@router.get("/applications/{application_no}/memo")
async def export_memo(application_no: str, actor_id: str = Query(...),
                      format: str = Query("pdf", pattern="^(pdf|docx)$")):
    from . import decisions
    actor = await run_in_threadpool(decisions._actor, actor_id)
    if actor["role_code"] not in ("officer", "head"):
        raise HTTPException(403, "The credit memo is an internal document for officers and the Head Manager.")
    model = await run_in_threadpool(build_model, application_no)
    if actor["role_code"] == "officer":
        app = await run_in_threadpool(decisions._load, application_no)
        if not decisions._same(app["assigned_officer_id"], actor["user_id"]):
            raise HTTPException(403, f"This application is assigned to {app['officer_name'] or 'another officer'}.")
    data = await run_in_threadpool(render_pdf if format == "pdf" else render_docx, model)
    name = f"CreditMemo_{application_no}_{datetime.now():%Y%m%d}.{format}"
    return Response(content=data, media_type=MEDIA[format],
                    headers={"Content-Disposition": f'attachment; filename="{name}"', "Cache-Control": "no-store"})
